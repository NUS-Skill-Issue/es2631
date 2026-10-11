"""Focused helper regressions. Set TMPDIR outside the repository for fixtures.

Run with PYTHONDONTWRITEBYTECODE=1 python3 -m unittest discover -s tests -v.
PPTX/PDF integration checks are optional when the reader dependencies exist.
"""
import contextlib
import importlib.util
import io
import os
from pathlib import Path
import runpy
import subprocess
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
CHECK = runpy.run_path(str(SCRIPTS / "check_essay.py"))
EXTRACT = runpy.run_path(str(SCRIPTS / "extract_deck.py"))


def slide(**changes):
    result = dict(n=1, title="Example", body=[], tables=[], charts=[],
                  pictures=0, notes="", warnings=[])
    result.update(changes)
    return result


class CheckerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="es2631-check-")
        self.addCleanup(self.temp.cleanup)
        self.essay = Path(self.temp.name) / "essay.txt"

    def run_check(self, text, *args):
        self.essay.write_text(text, encoding="utf-8")
        return subprocess.run([sys.executable, str(SCRIPTS / "check_essay.py"),
                               str(self.essay), *args], text=True, capture_output=True,
                              env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"}, timeout=15)

    def test_whitespace_blank_lines_and_short_prose_are_preserved(self):
        self.assertEqual(CHECK["split_paragraphs"]("First\n \nSecond\n\t\nThird"),
                         ["First", "Second", "Third"])

    def test_explicit_markdown_headings_and_references_are_excluded(self):
        self.assertEqual(CHECK["split_paragraphs"]("# Heading\n\nProse.\n\n## References\n\nCitation"),
                         ["Prose."])

    def test_default_does_not_impose_structure_keywords_or_elements(self):
        result = self.run_check("One short paragraph.")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertNotIn("mapping expects", result.stdout)
        self.assertNotIn("KEYWORD OBSERVATIONS", result.stdout)
        self.assertNotIn("MISSING", result.stdout)
        self.assertNotIn("STYLE OBSERVATIONS", result.stdout)

    def test_mapping_observes_keywords_without_quota_or_pass_score(self):
        result = self.run_check("Significance means the most important core.",
                                "--standards", "significance", "--paragraphs-only")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("Matches can overlap", result.stdout)
        self.assertIn("There is no keyword quota", result.stdout)
        self.assertNotIn("[PASS]", result.stdout)
        self.assertNotIn("/2 needed", result.stdout)

    def test_limit_is_a_hard_maximum_and_has_no_minimum(self):
        over = self.run_check(" ".join(["word"] * 101), "--limit", "100")
        under = self.run_check(" ".join(["word"] * 10), "--limit", "100")
        self.assertEqual(over.returncode, 1, over.stderr)
        self.assertIn("Exceeds the supplied maximum by 1", over.stdout)
        self.assertEqual(under.returncode, 0, under.stderr)
        self.assertNotIn("[LIMIT]", under.stdout)

    def test_element_requirement_is_explicit(self):
        result = self.run_check("The purpose is clear.", "--require-all-elements")
        self.assertEqual(result.returncode, 1, result.stderr)
        self.assertIn("[MATCH] purpose", result.stdout)
        self.assertIn("[MISSING] assumptions", result.stdout)
        self.assertIn("does not establish", result.stdout)

    def test_style_scan_does_not_claim_authorship_detection(self):
        result = self.run_check("This is a robust proposal.", "--style")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("do not detect human or AI authorship", result.stdout)
        self.assertIn("Words to review in context", result.stdout)

    def test_generated_metadata_numbers_are_not_source_evidence(self):
        source = ("# Deck extract: deck2.pptx\n7 slides.\n## Slide 7: Topic\n"
                  "[chart] plot 1 series 'cost': 20\n[table 3]\n[4 image(s) on this slide]\n"
                  "## Slides to inspect as images\n5, 6\n")
        tokens = CHECK["numeric_tokens"](CHECK["source_content"](source))
        self.assertEqual(tokens, {"20"})

    def test_numeric_matching_ignores_terminal_period_and_retains_sign(self):
        self.assertEqual(CHECK["numeric_tokens"]("20. −5 1,200.5"), {"20", "-5", "1200.5"})
        report = "\n".join(CHECK["claims"](["The cost was 20 dollars."], "The cost was 20."))
        self.assertNotIn("not found", report)
        self.assertIn("does not verify", report)

    def test_single_quotes_and_common_attribution_are_reviewed(self):
        report = "\n".join(CHECK["claims"](["IBM claims that over half improved. Patel states 'everyone improved after treatment'."], ""))
        self.assertIn("[REVIEW]", report)

    def test_invalid_limit_has_actionable_error(self):
        result = self.run_check("Text.", "--limit", "-1")
        self.assertEqual(result.returncode, 2)
        self.assertIn("positive integer", result.stderr)
        self.assertNotIn("Traceback", result.stderr)


class ExtractorTests(unittest.TestCase):
    def test_chart_extraction_keeps_all_values_and_plots(self):
        chart = SimpleNamespace(has_title=False, plots=[
            SimpleNamespace(categories=list(range(25)), series=[SimpleNamespace(name="first", values=list(range(25)))]),
            SimpleNamespace(categories=["later"], series=[SimpleNamespace(name="second", values=[9876])])])
        lines = EXTRACT["_chart_lines"](chart)
        self.assertTrue(any("22, 23, 24" in item for item in lines))
        self.assertTrue(any("9876" in item and "plot 2" in item for item in lines))

    def test_soft_line_break_does_not_join_words(self):
        frame = SimpleNamespace(paragraphs=[SimpleNamespace(text="first\vsecond", level=0)])
        result = "\n".join(EXTRACT["_paragraph_lines"](frame))
        self.assertNotIn("firstsecond", result)
        self.assertIn("first\n", result)

    def test_pdf_merging_preserves_negative_chart_labels(self):
        result = EXTRACT["_merge_pdf_lines"](["• Costs", "-5", "-10", "A separate label"])
        self.assertEqual(result, ["Costs", "-5", "-10", "A separate label"])

    def test_visual_content_is_flagged_despite_substantial_prose(self):
        for changes in ({"pictures": 1}, {"charts": ["[chart] WARNING unreadable"]},
                        {"warnings": ["Table extraction failed"]}):
            with self.subTest(changes=changes):
                _, flagged = EXTRACT["render"]([slide(body=["word " * 100], **changes)], "deck.pptx")
                self.assertEqual(flagged, [1])

    def test_speaker_notes_do_not_hide_sparse_slide(self):
        _, flagged = EXTRACT["render"]([slide(notes="word " * 100)], "deck.pptx")
        self.assertEqual(flagged, [1])

    def test_no_flags_does_not_claim_complete_extraction(self):
        rendered, flagged = EXTRACT["render"]([slide(body=["word " * 100])], "deck.pptx")
        self.assertEqual(flagged, [])
        self.assertIn("does not establish complete extraction", rendered)

    def run_pdf_fallback(self, result):
        with patch.dict(sys.modules, {"pdfplumber": None, "pypdf": None}), \
                patch("shutil.which", return_value="/mock/pdftotext"), \
                patch("subprocess.run", return_value=result):
            return EXTRACT["extract_pdf"]("fixture.pdf")

    def test_failed_pdf_converter_is_an_error(self):
        with self.assertRaisesRegex(RuntimeError, "pdftotext failed"):
            self.run_pdf_fallback(SimpleNamespace(returncode=1, stdout="", stderr="invalid input"))

    def test_pdf_fallback_preserves_blank_pages_and_source_positions(self):
        result = self.run_pdf_fallback(SimpleNamespace(returncode=0, stdout="\fSecond page\f", stderr=""))
        self.assertEqual([page["n"] for page in result], [1, 2])
        self.assertEqual(result[0]["title"], "")
        self.assertEqual(result[1]["title"], "Second page")
        self.assertEqual(EXTRACT["render"](result, "fixture.pdf")[1], [1, 2])

    def test_source_and_existing_output_are_protected(self):
        with tempfile.TemporaryDirectory(prefix="es2631-write-") as folder:
            source = Path(folder) / "source.pdf"
            output = Path(folder) / "output.md"
            source.write_bytes(b"original source")
            output.write_text("existing output", encoding="utf-8")
            for target, extra in ((source, ["--force"]), (output, [])):
                result = subprocess.run([sys.executable, str(SCRIPTS / "extract_deck.py"),
                                         str(source), "--out", str(target), *extra],
                                        text=True, capture_output=True, timeout=15,
                                        env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"})
                self.assertEqual(result.returncode, 2, result.stderr)
            self.assertEqual(source.read_bytes(), b"original source")
            self.assertEqual(output.read_text(encoding="utf-8"), "existing output")

    @unittest.skipUnless(importlib.util.find_spec("pptx"), "python-pptx is optional and unavailable")
    def test_actual_pptx_chart_and_soft_break(self):
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches
        with tempfile.TemporaryDirectory(prefix="es2631-pptx-") as folder:
            presentation = Presentation()
            page = presentation.slides.add_slide(presentation.slide_layouts[6])
            textbox = page.shapes.add_textbox(Inches(0), Inches(0), Inches(5), Inches(1))
            textbox.text_frame.paragraphs[0].text = "first\vsecond"
            data = CategoryChartData()
            data.categories = [f"c{n}" for n in range(25)]
            data.add_series("All values", list(range(25)))
            page.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0), Inches(1), Inches(5), Inches(4), data)
            path = Path(folder) / "sample.pptx"
            presentation.save(path)
            pages = EXTRACT["extract_pptx"](str(path))
            self.assertIn("first\n", "\n".join(pages[0]["body"]))
            self.assertTrue(any("24" in item and "series" in item for item in pages[0]["charts"]))

    @unittest.skipUnless(importlib.util.find_spec("pypdf"), "pypdf is optional and unavailable")
    def test_actual_blank_pdf_page_is_preserved(self):
        from pypdf import PdfWriter
        with tempfile.TemporaryDirectory(prefix="es2631-pdf-") as folder:
            path = Path(folder) / "blank.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=612, height=792)
            writer.write(path)
            pages = EXTRACT["extract_pdf"](str(path))
            self.assertEqual(len(pages), 1)
            self.assertEqual(EXTRACT["render"](pages, str(path))[1], [1])


if __name__ == "__main__":
    unittest.main()
